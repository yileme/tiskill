#!/usr/bin/env python3
"""Offline reading records. Run --help; all writes require prior user authorization."""
import argparse
from contextlib import contextmanager
from datetime import date
import fcntl
import hashlib
from html.parser import HTMLParser
import importlib.util
import json
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import tempfile
from urllib.parse import urlsplit, unquote
import uuid
import zipfile
import xml.etree.ElementTree as ET

MARK = "<!-- tis-reading -->\n"
STAGES = {"imported", "reading", "selected", "practising", "review"}

def digest(b): return hashlib.sha256(b).hexdigest()
def require(ok, message):
    if not ok: raise ValueError(message)
def ident(v):
    require(isinstance(v,str) and re.fullmatch(r"[A-Za-z0-9_-]{1,100}",v), "Unsafe ID")
    return v
def text(v):
    require(isinstance(v,str) and 0 < len(v.strip()) <= 100000, "Expected nonempty bounded text")
    return v

def folder(root, name):
    p = root / name
    require(not p.is_symlink(), "Symlink folder refused")
    p.mkdir(exist_ok=True)
    require(p.resolve().parent == root.resolve(), "Folder escapes root")
    return p

@contextmanager
def lock(root):
    p=root/'.tiskill-library.lock'
    require(not p.is_symlink(), "Symlink lock refused")
    with p.open('a') as f:
        fcntl.flock(f,fcntl.LOCK_EX)
        yield

def atomic_new(path, content):
    fd,tmp=tempfile.mkstemp(prefix='.reading-',dir=path.parent)
    try:
        with os.fdopen(fd,'w',encoding='utf-8') as f:
            f.write(content); f.flush(); os.fsync(f.fileno())
        os.link(tmp,path)  # exclusive publication; never replace existing notes
    finally:
        os.unlink(tmp)

def load_record(p):
    require(not p.is_symlink(), "Symlink record refused")
    t=p.read_text()
    if MARK not in t: return None
    return json.loads(t.rsplit(MARK,1)[1])

def all_records(root):
    out=[]
    for name in ['07-阅读素材','03-练习复盘']:
        d=root/name
        require(not d.is_symlink(), "Symlink folder refused")
        for p in d.glob('*.md'):
            v=load_record(p)
            if v: out.append((p,v))
    return out

def get_source(root, sid):
    found=[(p,v) for p,v in all_records(root) if v.get('kind')=='source' and v['source_id']==sid]
    require(len(found)==1,"Missing or duplicate canonical source")
    p,s=found[0]
    require(s.get('source_hash') and s.get('locator') and s.get('revision')==1 and Path(s.get('original_path','')).is_absolute(), 'Invalid canonical source')
    require(digest(s['content'].encode())==s['content_hash'],'Source snapshot changed')
    require(s.get('verification') in ('unverified','text_read'), 'Invalid source verification')
    if s['verification']=='text_read':
        require(s.get('actual_read') is True and s.get('user_confirmed') is True and s.get('checker') and not s.get('ocr_unchecked'), 'Canonical reading proof missing')
    return p,s

def store(root, rec, name, stem):
    p=folder(root,name)/(stem+'.md')
    body='# '+rec.get('title',rec['kind'])+'\n\n[[首页]]\n\n'
    if rec['kind']=='source':
        for i,para in enumerate(rec['content'].split('\n\n'),1):
            body+=para+'\n^'+rec['source_id']+'-p'+str(i)+'\n\n'
    else:
        body+='```json\n'+json.dumps(rec,ensure_ascii=False,indent=2)+'\n```\n\n'
    atomic_new(p,body+MARK+json.dumps(rec,ensure_ascii=False,indent=2)+'\n')
    return str(p)

def span(value, total):
    require(isinstance(value,str) and re.fullmatch(r'[0-9]+(-[0-9]+)?',value),'Expected inclusive 1-based range')
    a,*rest=map(int,value.split('-')); b=rest[0] if rest else a
    require(1<=a<=b<=total and b-a<100,'Range out of bounds or over 100 units')
    return a,b

class Plain(HTMLParser):
    def __init__(self): super().__init__(convert_charrefs=True); self.parts=[]; self.skip=0
    def handle_starttag(self,t,a):
        if t in ('script','style'): self.skip+=1
        if t in ('p','div','br','h1','h2','li'): self.parts.append('\n')
    def handle_endtag(self,t):
        if t in ('script','style'): self.skip=max(0,self.skip-1)
    def handle_data(self,d):
        if not self.skip:self.parts.append(d)

def xml(b):
    require(not re.search(br'<!\s*(DOCTYPE|ENTITY)',b,re.I),'DTD/entities refused')
    return ET.fromstring(b)

def zipname(s):
    s=unquote(s)
    require(not s.startswith('/') and '\\' not in s and '..' not in PurePosixPath(s).parts,'Unsafe ZIP path')
    return s

def convert(d):
    p=Path(d['path']).expanduser().resolve(strict=True)
    require(p.is_file() and p.stat().st_size<=100*1024*1024,'Input exceeds 100 MB bound')
    raw=p.read_bytes(); ext=p.suffix.lower()
    if ext in ('.md','.txt'):
        lines=raw.decode('utf-8').splitlines()
        if d.get('lines'): a,b=span(d['lines'],len(lines))
        else:
            require(d.get('bounded_document') is True and len(raw)<=20000,'Select lines or explicitly accept bounded document <=20KB')
            a,b=1,len(lines)
        content='\n'.join(lines[a-1:b]); locator=f'lines {a}-{b}'
    elif ext=='.pdf':
        require(d.get('pages'),'Select PDF pages')
        require(importlib.util.find_spec('pypdf') is not None,'Install pypdf for PDF extraction')
        from pypdf import PdfReader
        pdf=PdfReader(p); a,b=span(d['pages'],len(pdf.pages))
        pages=[page.extract_text() or '' for page in pdf.pages[a-1:b]]
        require(all(v.strip() for v in pages),'Empty/scanned page: OCR deferred; supply checked transcription separately')
        content='\n\n'.join(f'[page {a+i}]\n'+v for i,v in enumerate(pages));locator=f'PDF pages {a}-{b}'
    elif ext=='.epub':
        require(type(d.get('chapter')) is int,'Select 1-based EPUB spine chapter')
        with zipfile.ZipFile(p) as z:
            infos=z.infolist(); names=[i.filename for i in infos]
            require(len(names)==len(set(names)) and len(names)<=10000,'Duplicate/excess ZIP members')
            for i in infos:
                zipname(i.filename)
                require(i.file_size<=10*1024*1024,'Oversize ZIP member')
                require((i.external_attr>>16)&0o170000 != 0o120000,'ZIP symlink refused')
            require(sum(i.file_size for i in infos)<=100*1024*1024,'ZIP expansion exceeds bound')
            container=xml(z.read('META-INF/container.xml'))
            opf=zipname(next(e.attrib['full-path'] for e in container.iter() if e.tag.endswith('rootfile')))
            tree=xml(z.read(opf)); base=PurePosixPath(opf).parent
            manifest={e.attrib['id']:e.attrib['href'] for e in tree.iter() if e.tag.endswith('}item') or e.tag=='item'}
            spine=[e.attrib['idref'] for e in tree.iter() if e.tag.endswith('}itemref') or e.tag=='itemref']
            n=d['chapter'];require(1<=n<=len(spine),'Chapter out of bounds')
            href=manifest[spine[n-1]];require(not urlsplit(href).scheme,'External chapter refused')
            entry=zipname(str(base/href.split('#')[0]));data=z.read(entry)
            xml(data); parser=Plain();parser.feed(data.decode('utf-8'));content=''.join(parser.parts)
            locator=f'EPUB spine {n}: {href}'
    else: raise ValueError('Supported: MD, TXT, clipped web MD, PDF, EPUB')
    text(content)
    return p,digest(raw),content,locator

def source_add(root,d):
    sid=ident(d['source_id']);require(not any(v.get('source_id')==sid and v['kind']=='source' for _,v in all_records(root)),'Source ID collision')
    p,h,c,loc=convert(d)
    url=d.get('url')
    if url:
        u=urlsplit(url);require(u.scheme in ('http','https') and u.hostname and not u.username and not u.password and not re.search(r'[\s\x00-\x1f]',url),'Invalid URL metadata')
    verified=d.get('verification','unverified');require(verified in ('unverified','text_read'),'Unsupported verification')
    if verified=='text_read':
        require(d.get('actual_read') is True and d.get('user_confirmed') is True and not d.get('ocr_unchecked') and d.get('checker'),'Actual read and user confirmation/checker required; unchecked OCR cannot verify')
    rec=dict(kind='source',source_id=sid,title=text(d['title']),revision=1,source_hash=h,content_hash=digest(c.encode()),original_path=str(p),content=c,locator=loc,url=url,verification=verified,checker=d.get('checker'),actual_read=d.get('actual_read') is True,user_confirmed=d.get('user_confirmed') is True,ocr_unchecked=bool(d.get('ocr_unchecked')),stage='imported')
    return store(root,rec,'07-阅读素材',sid+'-source-r1')

def study(root,d):
    require('snapshot' not in d,'Input snapshot refused')
    p,s=get_source(root,d['source_id']);stage=d.get('stage','reading');require(stage in STAGES,'Invalid stage')
    attempt=ident(d['attempt_id']);require(not any(v.get('attempt_id')==attempt for _,v in all_records(root)),'Attempt collision')
    for k in ('summary','response','self_review'): text(d[k])
    rec=dict(d,kind='study',snapshot=s,source_link=str(p.relative_to(root)),date=d.get('date',date.today().isoformat()),stage=stage)
    date.fromisoformat(rec['date'])
    return store(root,rec,'03-练习复盘',attempt)

def card(root,d):
    require('snapshot' not in d,'Input snapshot refused')
    _,s=get_source(root,d['source_id']);cid=ident(d['card_id']);revision=d.get('revision',1)
    prev=[v for _,v in all_records(root) if v['kind']=='card' and v['card_id']==cid]
    require(type(revision) is int and revision==max([v['revision'] for v in prev],default=0)+1,'Card ID/revision collision or gap')
    if prev: require(d.get('revise') is True,'Explicit revise required')
    require(d.get('locator')==s['locator'],'Locator must match source')
    for k in ('question','answer','type'):text(d[k])
    status=d.get('status','draft');require(status in ('draft','confirmed'),'Invalid card status')
    if status=='confirmed':require(d.get('user_confirmed') is True,'User confirmation required')
    origin=d.get('origin');require(origin in ('source_quote','model_generated','learner'),'Mark card origin')
    if origin=='source_quote':require(text(d.get('quote')) in s['content'],'Quote does not match source')
    if d.get('due'):date.fromisoformat(d['due'])
    require(isinstance(d.get('hint',''),str),'Invalid hint')
    rec=dict(d,kind='card',revision=revision,status=status,snapshot=s,export_state='pending')
    return store(root,rec,'07-阅读素材',cid+'-card-r'+str(revision))

def latest(root):
    out={}
    for _,v in all_records(root):
        if v['kind']=='card' and v['revision']>out.get(v['card_id'],{}).get('revision',0):out[v['card_id']]=v
    return list(out.values())

def manifests(root):
    p=root/'07-阅读素材'/'exports';require(not p.is_symlink(),'Symlink exports refused')
    out=[]
    for b in p.glob('*'):
        require(not b.is_symlink(),'Symlink batch refused')
        if b.name.startswith('.'):continue
        m=b/'manifest.json';require(not m.is_symlink(),'Symlink manifest refused')
        if m.exists():out.append(json.loads(m.read_text()))
    return out

def export(root, include=False):
    cards=latest(root); old={(c['card_id'],c['revision']) for m in manifests(root) for c in m['cards']}
    changed=[c['card_id'] for c in cards if any(i==c['card_id'] for i,r in old) and (c['card_id'],c['revision']) not in old]
    selected=[c for c in cards if c['status']=='confirmed' and (c['card_id'],c['revision']) not in old and (include or c['card_id'] not in changed)]
    if not selected:return dict(exported=0,changed_previously_exported=changed)
    require(importlib.util.find_spec('openpyxl') is not None,'Install openpyxl for XLSX export')
    from openpyxl import Workbook
    out=folder(folder(root,'07-阅读素材'),'exports');batch='batch-'+uuid.uuid4().hex
    tmp=Path(tempfile.mkdtemp(prefix='.pending-',dir=out))
    try:
        files=[]
        for hint in (False,True):
            group=[c for c in selected if bool(c.get('hint'))==hint]
            if not group:continue
            rows=[]
            for c in group:
                s=c['snapshot'];citation=f" [{s['title']} · {s['locator']} · {s['source_id']}]"
                row=[c['question'],c['answer']+(citation if c.get('citation',True) else '')]
                if hint:row.append(c['hint'])
                # Exports never include local paths from canonical metadata.
                rows.append([' '.join(v.split()) for v in row])
            name='hint' if hint else 'regular'; tsv=tmp/(name+'.tsv');tsv.write_text(''.join('\t'.join(row)+'\n' for row in rows))
            wb=Workbook();ws=wb.active
            for row in rows:
                ws.append(row)
                for cell in ws[ws.max_row]:cell.data_type='s'
            xp=tmp/(name+'.xlsx');wb.save(xp)
            require(xp.stat().st_size<=10*1024*1024,'XLSX exceeds Qingk 10MB limit')
            files += [tsv.name,xp.name]
        manifest=dict(batch_id=batch,state='exported_not_imported',cards=[dict(card_id=c['card_id'],revision=c['revision']) for c in selected],files={n:digest((tmp/n).read_bytes()) for n in files})
        (tmp/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
        os.rename(tmp,out/batch)
        return dict(exported=len(selected),batch=str(out/batch),changed_previously_exported=changed)
    finally:
        if tmp.exists():shutil.rmtree(tmp)

def imported(root,d):
    require(d.get('user_confirmed') is True,'Manual user import confirmation required')
    batch=ident(d['batch_id']);require(any(m['batch_id']==batch for m in manifests(root)),'Unknown batch')
    rec=dict(kind='import_ack',batch_id=batch,user_confirmed=True,date=date.today().isoformat())
    return store(root,rec,'07-阅读素材',batch+'-import-ack')

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--config',type=Path,default=Path.home()/'.tiskill/library.json')
    p.add_argument('command',choices=['sources','study','card','listdue','export','mark-imported'])
    p.add_argument('--input',type=Path);p.add_argument('--confirm',action='store_true');p.add_argument('--include-revisions',action='store_true');p.add_argument('--on',default=date.today().isoformat())
    a=p.parse_args()
    try:
        root=Path(json.loads(a.config.read_text())['path']).expanduser();require(root.is_absolute() and root.is_dir(),'Existing absolute library root required');root=root.resolve()
        if a.command=='listdue':
            day=date.fromisoformat(a.on);result=[dict(card_id=c['card_id'],revision=c['revision'],status=c['status'],due=c['due']) for c in latest(root) if c.get('due') and date.fromisoformat(c['due'])<=day]
        else:
            require(a.confirm,'Writes require --confirm after user authorization')
            with lock(root):
                if a.command=='export':result=export(root,a.include_revisions)
                else:
                    require(a.input is not None,'--input JSON required');d=json.loads(a.input.read_text());result={'sources':source_add,'study':study,'card':card,'mark-imported':imported}[a.command](root,d)
        print(json.dumps(result,ensure_ascii=False))
    except (ValueError,KeyError,OSError,ET.ParseError,zipfile.BadZipFile,StopIteration) as e:p.exit(2,str(e)+'\n')
if __name__=='__main__':main()
