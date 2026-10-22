"""Isolated synthetic fixtures; never use private reading material."""
import importlib.util
import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

SCRIPT=Path(__file__).resolve().parents[1]/'skills/tis-reading/scripts/reading.py'
spec=importlib.util.spec_from_file_location('reading',SCRIPT)
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
from tools import library

class Reading(unittest.TestCase):
    def setUp(self):
        self.t=tempfile.TemporaryDirectory();self.addCleanup(self.t.cleanup);self.root=Path(self.t.name)
        self.src=self.root/'source.md';self.src.write_text('中文 “quoted”\n\nA useful chunk.\n')
        self.original=self.src.read_bytes()
        self.source={'source_id':'s1','title':'Synthetic 中文','path':str(self.src),'bounded_document':True}
        m.source_add(self.root,self.source)
    def card(self,**kw):
        d=dict(card_id='c1',source_id='s1',locator='lines 1-3',type='chunk',question='=1+1\t中文',answer='@quoted\nanswer',origin='model_generated',status='confirmed',user_confirmed=True,due='2026-11-02');d.update(kw);return d
    def test_bounds_snapshot_and_proof(self):
        self.assertEqual(self.src.read_bytes(),self.original)
        with self.assertRaises(ValueError):m.source_add(self.root,dict(self.source,source_id='s2',bounded_document=False))
        with self.assertRaises(ValueError):m.source_add(self.root,dict(self.source,source_id='s2',verification='text_read'))
        with self.assertRaises(ValueError):m.source_add(self.root,dict(self.source,source_id='s2',verification='text_read',actual_read=True,user_confirmed=True,checker='user',ocr_unchecked=True))
        d=dict(attempt_id='a1',source_id='s1',summary='我的总结',response='my answer',self_review='pending',stage='practising')
        p=Path(m.study(self.root,d));self.assertEqual(m.load_record(p)['snapshot']['content'],'中文 “quoted”\n\nA useful chunk.')
        with self.assertRaises(ValueError):m.study(self.root,d)
        m.study(self.root,dict(d,attempt_id='a2',response='second attempt'))
        self.assertEqual(m.load_record(p)['response'],'my answer')
        e=dict(type='expression',title='Local example',expression='chunk',context='context',explanation='explanation',quotation='A useful chunk.',localSourceId='s1',verification='text_read',userConfirmed=True)
        with self.assertRaises(ValueError):library.save(self.root,e,True)
        m.source_add(self.root,dict(self.source,source_id='s2',verification='text_read',actual_read=True,user_confirmed=True,checker='user'))
        path=library.save(self.root,dict(e,localSourceId='s2'),True)
        self.assertEqual(library.payload(path)['localSourceSnapshot']['source_id'],'s2')
    def test_export_and_revisions(self):
        from openpyxl import load_workbook
        m.card(self.root,self.card());m.card(self.root,self.card(card_id='hint',hint='提示'))
        m.card(self.root,self.card(card_id='pending',status='draft'))
        with self.assertRaises(ValueError):m.card(self.root,self.card())
        with self.assertRaises(ValueError):m.card(self.root,self.card(card_id='bad',origin='source_quote',quote='invented'))
        x=m.export(self.root);self.assertEqual(x['exported'],2);b=Path(x['batch'])
        for name,cols in [('regular',2),('hint',3)]:
            ws=load_workbook(b/(name+'.xlsx')).active
            self.assertEqual((ws.max_row,ws.max_column),(1,cols))
            self.assertEqual(ws['A1'].data_type,'s');self.assertTrue(ws['A1'].value.startswith('='))
            rows=(b/(name+'.tsv')).read_text().splitlines();self.assertEqual(len(rows),1);self.assertEqual(len(rows[0].split('\t')),cols)
            self.assertNotIn(str(self.root),(b/(name+'.tsv')).read_text())
        self.assertEqual(m.export(self.root)['exported'],0)
        m.card(self.root,self.card(revision=2,revise=True,answer='changed'))
        self.assertEqual(m.export(self.root)['changed_previously_exported'],['c1'])
        self.assertEqual(m.export(self.root,True)['exported'],1)
        with self.assertRaises(ValueError):m.imported(self.root,dict(batch_id=b.name))
        m.imported(self.root,dict(batch_id=b.name,user_confirmed=True))
        self.assertEqual(next(c for c in m.latest(self.root) if c['card_id']=='pending')['status'],'draft')
    def test_transient_failure_and_due(self):
        m.card(self.root,self.card())
        with patch.object(m.os,'rename',side_effect=OSError('transient')):
            with self.assertRaises(OSError):m.export(self.root)
        self.assertEqual(m.manifests(self.root),[])
        self.assertEqual(m.export(self.root)['exported'],1)
        config=self.root/'config.json';config.write_text(json.dumps(dict(path=str(self.root))))
        before={str(p):p.read_bytes() for p in self.root.rglob('*.md')}
        for day,count in [('2026-11-01',0),('2026-11-02',1),('2026-11-03',1)]:
            run=subprocess.run([sys.executable,str(SCRIPT),'--config',str(config),'listdue','--on',day],capture_output=True,text=True,check=True)
            self.assertEqual(len(json.loads(run.stdout)),count)
        self.assertEqual(before,{str(p):p.read_bytes() for p in self.root.rglob('*.md')})
    def epub(self,p,unsafe=False):
        with zipfile.ZipFile(p,'w') as z:
            z.writestr('META-INF/container.xml','<container><rootfiles><rootfile full-path="OPS/book.opf"/></rootfiles></container>')
            z.writestr('OPS/book.opf','<package><manifest><item id="a" href="a.xhtml"/><item id="b" href="b.xhtml"/></manifest><spine><itemref idref="b"/><itemref idref="a"/></spine></package>')
            z.writestr('OPS/a.xhtml','<html><body><p>First in ZIP</p></body></html>')
            z.writestr('OPS/b.xhtml','<html><body><p>中文 spine first</p></body></html>')
            if unsafe:z.writestr('../escape','bad')
    def test_epub_and_pdf(self):
        ep=self.root/'book.epub';self.epub(ep);raw=ep.read_bytes()
        _,_,content,loc=m.convert(dict(path=str(ep),chapter=1));self.assertIn('中文 spine first',content);self.assertIn('spine 1',loc);self.assertEqual(ep.read_bytes(),raw)
        self.epub(ep,True)
        with self.assertRaises(ValueError):m.convert(dict(path=str(ep),chapter=1))
        with self.assertRaises(ValueError):m.xml(b'<!DOCTYPE x [<!ENTITY e "bad">]><x/>')
        import fitz
        pdf=self.root/'tiny.pdf';doc=fitz.open()
        for value in ['first page','selected page']:doc.new_page().insert_text((50,50),value)
        doc.save(pdf);doc.close();raw=pdf.read_bytes()
        _,_,content,loc=m.convert(dict(path=str(pdf),pages='2'));self.assertIn('selected page',content);self.assertNotIn('first page',content);self.assertEqual(pdf.read_bytes(),raw)
        with self.assertRaises(ValueError):m.convert(dict(path=str(pdf)))
    def test_standalone(self):
        dst=self.root/'standalone';shutil.copytree(SCRIPT.parent.parent,dst)
        run=subprocess.run([sys.executable,str(dst/'scripts/reading.py'),'--help'],capture_output=True,text=True)
        self.assertEqual(run.returncode,0)

if __name__=='__main__':unittest.main()
